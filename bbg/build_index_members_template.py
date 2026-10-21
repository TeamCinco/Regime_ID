from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT_PATH = Path(__file__).resolve().parent / "bbg_index_members.xlsx"
INDEX_SLOTS = 6
SNAPSHOT_COUNT = 48
FIRST_INPUT_ROW = 11
EXAMPLE_INDEXES = [
    ("S5HLTH Index", "Health care"),
    ("S5ENRS Index", "Energy"),
    ("S5INSU Index", "Insurance"),
    ("S5CONS Index", "Staples"),
]
FIRST_SNAPSHOT_DATE = date(2015, 1, 1)
MONTHS_BETWEEN_SNAPSHOTS = 3
MEMBERS_FIELD = "INDX_MWEIGHT_HIST"
DATE_OVERRIDE = "END_DATE_OVERRIDE"
BLOCK_WIDTH = 3

BASE_FONT = Font(name="Arial", size=10)
BOLD_FONT = Font(name="Arial", size=10, bold=True)
TITLE_FONT = Font(name="Arial", size=14, bold=True)
INPUT_FONT = Font(name="Arial", size=10, color="0000FF")
LINK_FONT = Font(name="Arial", size=10, color="008000")
INDEX_LINK_FONT = Font(name="Arial", size=11, bold=True, color="008000")
INPUT_FILL = PatternFill("solid", start_color="FFFF00")
HEADER_FILL = PatternFill("solid", start_color="D9E1F2")


def build_workbook():
    workbook = Workbook()
    write_inputs_sheet(workbook.active)
    write_members_sheet(workbook.create_sheet("Members"))
    write_notes_sheet(workbook.create_sheet("Notes"))
    workbook.save(OUTPUT_PATH)


def write_inputs_sheet(sheet):
    sheet.title = "Inputs"
    write_cell(sheet, "A1", "Bloomberg point-in-time index members", TITLE_FONT)
    write_cell(sheet, "A3", "First snapshot date", BOLD_FONT)
    write_input(sheet, "B3", FIRST_SNAPSHOT_DATE).number_format = "m/d/yyyy"
    write_cell(sheet, "A4", "Months between snapshots", BOLD_FONT)
    write_input(sheet, "B4", MONTHS_BETWEEN_SNAPSHOTS)
    write_legend(sheet)
    write_snapshot_dates(sheet)
    write_index_table(sheet)
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["D"].width = 18
    sheet.column_dimensions["E"].width = 16
    sheet.column_dimensions["F"].width = 18


def write_legend(sheet):
    lines = [
        "Type Bloomberg index tickers and a peer group name in D11:E16 (yellow). Rows 11-14 are examples.",
        "Set the first date and spacing in B3:B4, then Bloomberg tab > Refresh Workbook.",
        "Each index gets one member list per snapshot date on the Members tab (dates after today stay blank).",
    ]
    write_cell(sheet, "D3", "How to use", BOLD_FONT)
    for offset, line in enumerate(lines):
        write_cell(sheet, f"D{4 + offset}", line)


def write_snapshot_dates(sheet):
    write_cell(sheet, "A10", "Snapshot", BOLD_FONT).fill = HEADER_FILL
    write_cell(sheet, "B10", "Date", BOLD_FONT).fill = HEADER_FILL
    for snapshot in range(SNAPSHOT_COUNT):
        row = FIRST_INPUT_ROW + snapshot
        write_cell(sheet, f"A{row}", snapshot + 1)
        write_cell(sheet, f"B{row}", snapshot_date_formula(row)).number_format = "m/d/yyyy"


def snapshot_date_formula(row):
    if row == FIRST_INPUT_ROW:
        return "=$B$3"
    return f"=EDATE(B{row - 1},$B$4)"


def write_index_table(sheet):
    for offset, header in enumerate(["Index", "Peer group", "Members tab column"]):
        write_cell(sheet, f"{get_column_letter(4 + offset)}10", header, BOLD_FONT).fill = HEADER_FILL
    for slot in range(INDEX_SLOTS):
        write_index_row(sheet, slot)
    sheet.freeze_panes = "A11"


def write_index_row(sheet, slot):
    row = FIRST_INPUT_ROW + slot
    index_ticker, peer_group = EXAMPLE_INDEXES[slot] if slot < len(EXAMPLE_INDEXES) else (None, None)
    write_input(sheet, f"D{row}", index_ticker)
    write_input(sheet, f"E{row}", peer_group)
    write_cell(sheet, f"F{row}", get_column_letter(block_first_column(slot, 0)))


def write_members_sheet(sheet):
    for slot in range(INDEX_SLOTS):
        for snapshot in range(SNAPSHOT_COUNT):
            write_members_block(sheet, slot, snapshot)
    sheet.freeze_panes = "A5"


def write_members_block(sheet, slot, snapshot):
    first_column = block_first_column(slot, snapshot)
    index_cell = f"Inputs!$D${FIRST_INPUT_ROW + slot}"
    peer_group_cell = f"Inputs!$E${FIRST_INPUT_ROW + slot}"
    date_cell = f"Inputs!$B${FIRST_INPUT_ROW + snapshot}"
    write_cell(sheet, cell_at(first_column, 1), f"={index_cell}", INDEX_LINK_FONT)
    write_cell(sheet, cell_at(first_column + 1, 1), f"={peer_group_cell}", LINK_FONT)
    write_cell(sheet, cell_at(first_column, 2), f"={date_cell}", LINK_FONT).number_format = "m/d/yyyy"
    write_cell(sheet, cell_at(first_column, 3), "Member", BOLD_FONT).fill = HEADER_FILL
    write_cell(sheet, cell_at(first_column + 1, 3), "Weight", BOLD_FONT).fill = HEADER_FILL
    write_cell(sheet, cell_at(first_column, 4), bds_formula(index_cell, date_cell))
    set_block_widths(sheet, first_column)


def bds_formula(index_cell, date_cell):
    override = f'"{DATE_OVERRIDE}="&TEXT({date_cell},"yyyymmdd")'
    bds = f'BDS({index_cell},"{MEMBERS_FIELD}",{override})'
    return f'=IF(OR({index_cell}="",{date_cell}>TODAY()),"",{bds})'


def set_block_widths(sheet, first_column):
    sheet.column_dimensions[get_column_letter(first_column)].width = 12
    sheet.column_dimensions[get_column_letter(first_column + 1)].width = 10
    sheet.column_dimensions[get_column_letter(first_column + 2)].width = 3


def write_notes_sheet(sheet):
    notes = [
        f"{MEMBERS_FIELD} returns each member and its weight on the {DATE_OVERRIDE} date (yyyymmdd).",
        "Members come back as exchange tickers like 'UNH UN'; add ' Equity' to pull their data.",
        "Example indexes: S5HLTH = S&P 500 Health Care, S5ENRS = Energy, S5INSU = Insurance, "
        "S5CONS = Consumer Staples. Confirm tickers on the index's DES page.",
        "Swap in any index (SPX Index, RIY Index, S5INFT Index) to change the universe.",
        "If a block stays blank after refresh, remove the IF( wrapper so the cell is just =BDS(...).",
        "Next step: the union of all members becomes the stock list for the factor vol template.",
    ]
    write_cell(sheet, "A1", "Notes", BOLD_FONT)
    for offset, note in enumerate(notes):
        write_cell(sheet, f"A{2 + offset}", note)
    sheet.column_dimensions["A"].width = 110


def block_first_column(slot, snapshot):
    return 1 + (slot * SNAPSHOT_COUNT + snapshot) * BLOCK_WIDTH


def cell_at(column, row):
    return f"{get_column_letter(column)}{row}"


def write_cell(sheet, address, value, font=BASE_FONT):
    cell = sheet[address]
    cell.value = value
    cell.font = font
    return cell


def write_input(sheet, address, value):
    cell = write_cell(sheet, address, value, INPUT_FONT)
    cell.fill = INPUT_FILL
    return cell


if __name__ == "__main__":
    build_workbook()
