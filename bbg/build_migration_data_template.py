from datetime import date
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

TEMPLATE_FOLDER = Path(__file__).resolve().parent
EVENTS_CSV = TEMPLATE_FOLDER / "migration_events.csv"
OUTPUT_PATH = TEMPLATE_FOLDER / "bbg_migration_data.xlsx"
FIRST_TICKER_ROW = 11
MAX_SECURITIES = 1500
START_DATE = date(2014, 1, 1)
BENCHMARKS = [
    "SPX Index", "RLG Index", "RLV Index",
    "S5INFT Index", "S5HLTH Index", "S5FINL Index", "S5COND Index", "S5CONS Index", "S5ENRS Index",
    "S5INDU Index", "S5MATR Index", "S5UTIL Index", "S5RLST Index", "S5TELS Index",
]
FIELDS = [
    ("PX_LAST", "Last price"),
    ("CUR_MKT_CAP", "Market cap, in millions"),
    ("EV_TO_T12M_SALES", "Enterprise value to trailing 12-month sales"),
    ("BEST_PE_RATIO", "Forward P/E on consensus EPS"),
    ("BEST_EPS", "Consensus EPS (blended next 12 months), for revisions"),
    ("BEST_SALES", "Consensus sales (blended next 12 months), for revisions"),
    ("SALES_REV_TURN", "Trailing reported sales; stamped on period end, so the study lags it"),
    ("30DAY_IMPVOL_100.0%MNY_DF", "30-day at-the-money implied vol, in % points"),
]
STATIC_FIELD = ("GICS_SECTOR_NAME", "GICS sector")
BLOCK_WIDTH = len(FIELDS) + 2
BDH_OPTIONS = '"Dir=V","Per=W","Fill=P"'

BASE_FONT = Font(name="Arial", size=10)
BOLD_FONT = Font(name="Arial", size=10, bold=True)
TITLE_FONT = Font(name="Arial", size=14, bold=True)
INPUT_FONT = Font(name="Arial", size=10, color="0000FF")
LINK_FONT = Font(name="Arial", size=10, color="008000")
TICKER_LINK_FONT = Font(name="Arial", size=11, bold=True, color="008000")
INPUT_FILL = PatternFill("solid", start_color="FFFF00")
HEADER_FILL = PatternFill("solid", start_color="D9E1F2")


def build_workbook():
    securities = read_event_securities() + BENCHMARKS
    workbook = Workbook()
    write_inputs_sheet(workbook.active, securities)
    write_data_sheet(workbook.create_sheet("Data"), len(securities))
    write_fields_sheet(workbook.create_sheet("Fields"))
    workbook.save(OUTPUT_PATH)
    print(f"Saved {OUTPUT_PATH} with {len(securities)} securities; refresh it on the terminal")


def read_event_securities():
    if not EVENTS_CSV.exists():
        raise FileNotFoundError(f"{EVENTS_CSV} is missing; run Regime_ID/migration_universe.py first")
    securities = sorted(pd.read_csv(EVENTS_CSV)["security"].unique())
    if len(securities) + len(BENCHMARKS) > MAX_SECURITIES:
        raise ValueError(f"{len(securities)} migrated stocks is over the {MAX_SECURITIES} column budget")
    return securities


def write_inputs_sheet(sheet, securities):
    sheet.title = "Inputs"
    write_cell(sheet, "A1", "Bloomberg Growth -> Value migration data pull", TITLE_FONT)
    write_cell(sheet, "A3", "Start date", BOLD_FONT)
    write_input(sheet, "B3", START_DATE).number_format = "m/d/yyyy"
    write_cell(sheet, "A4", "End date (blank = today)", BOLD_FONT)
    write_input(sheet, "B4", None).number_format = "m/d/yyyy"
    write_cell(sheet, "A6", "Fields", BOLD_FONT)
    for offset, (field, _) in enumerate(FIELDS):
        write_input(sheet, cell_at(2 + offset, 6), field)
    write_cell(sheet, "A8", "Tickers come from migration_events.csv plus benchmarks. "
                            "Bloomberg tab > Refresh Workbook, then save.")
    write_ticker_table(sheet, securities)
    sheet.column_dimensions["A"].width = 24
    for offset in range(1, len(FIELDS) + 1):
        sheet.column_dimensions[get_column_letter(1 + offset)].width = 24


def write_ticker_table(sheet, securities):
    for offset, header in enumerate(["Bloomberg security", STATIC_FIELD[1], "Data tab column"]):
        write_cell(sheet, cell_at(1 + offset, 10), header, BOLD_FONT).fill = HEADER_FILL
    for slot, security in enumerate(securities):
        row = FIRST_TICKER_ROW + slot
        write_input(sheet, f"A{row}", security)
        write_cell(sheet, f"B{row}", f'=IF(A{row}="","",BDP(A{row},"{STATIC_FIELD[0]}"))')
        write_cell(sheet, f"C{row}", get_column_letter(block_first_column(slot)))
    sheet.freeze_panes = "A11"


def write_data_sheet(sheet, security_count):
    for slot in range(security_count):
        write_data_block(sheet, slot)
    sheet.freeze_panes = "A4"


def write_data_block(sheet, slot):
    first_column = block_first_column(slot)
    security_cell = f"Inputs!$A${FIRST_TICKER_ROW + slot}"
    write_cell(sheet, cell_at(first_column, 1), f"={security_cell}", TICKER_LINK_FONT)
    write_cell(sheet, cell_at(first_column, 2), "Date", BOLD_FONT).fill = HEADER_FILL
    for offset in range(len(FIELDS)):
        field_cell = f"Inputs!{get_column_letter(2 + offset)}$6"
        header = write_cell(sheet, cell_at(first_column + 1 + offset, 2), f"={field_cell}", LINK_FONT)
        header.fill = HEADER_FILL
        header.alignment = Alignment(wrap_text=True)
    write_cell(sheet, cell_at(first_column, 3), bdh_formula(security_cell))


def bdh_formula(security_cell):
    last_field_column = get_column_letter(1 + len(FIELDS))
    fields_range = f"Inputs!$B$6:${last_field_column}$6"
    end_date = 'IF(Inputs!$B$4="","",Inputs!$B$4)'
    bdh = f"BDH({security_cell},{fields_range},Inputs!$B$3,{end_date},{BDH_OPTIONS})"
    return f'=IF({security_cell}="","",{bdh})'


def write_fields_sheet(sheet):
    write_cell(sheet, "A1", "Field", BOLD_FONT).fill = HEADER_FILL
    write_cell(sheet, "B1", "What it is", BOLD_FONT).fill = HEADER_FILL
    for offset, (field, meaning) in enumerate(FIELDS):
        write_cell(sheet, f"A{2 + offset}", field)
        write_cell(sheet, f"B{2 + offset}", meaning)
    notes = [
        "Weekly data (Per=W) keeps hundreds of stocks inside Bloomberg's limits.",
        "Confirm field names in FLDS <GO>. Index blocks return #N/A for company fields; that is expected.",
        "If a block stays blank after refresh, remove the IF( wrapper so the cell is just =BDH(...).",
    ]
    for offset, note in enumerate(notes):
        write_cell(sheet, f"A{4 + len(FIELDS) + offset}", note)
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 70


def block_first_column(slot):
    return 1 + slot * BLOCK_WIDTH


def cell_at(column, row):
    return f"{get_column_letter(column)}{row}"


def write_cell(sheet, address, value, font=BASE_FONT):
    cell = sheet[address]
    cell.value = value
    cell.font = font
    return cell


def write_input(sheet, address, value):
    cell = write_cell(sheet, address, value, INPUT_FONT)
    cell.fill = INPUT_FILL
    return cell


if __name__ == "__main__":
    build_workbook()
