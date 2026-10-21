from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = Path(__file__).resolve().parent / "bbg_vol_data.xlsx"
TICKER_SLOTS = 30
FIRST_TICKER_ROW = 11
EXAMPLE_TICKERS = ["GLD", "SNDK", "SPY"]
DEFAULT_YELLOW_KEY = "US Equity"
START_DATE = date(2015, 1, 1)
FIELDS = [
    ("PX_LAST", "Last price"),
    ("VOLATILITY_30D", "30-day realized (historical) vol, in % points"),
    ("30DAY_IMPVOL_100.0%MNY_DF", "30-day at-the-money implied vol, in % points"),
    ("3MTH_IMPVOL_100.0%MNY_DF", "3-month at-the-money implied vol, in % points"),
    ("1M_PUT_IMP_VOL_25DELTA_DFLT", "1-month 25-delta put implied vol (skew), in % points"),
]
BLOCK_WIDTH = len(FIELDS) + 2
BDH_OPTIONS = '"Dir=V","Days=T","Fill=P"'

BASE_FONT = Font(name="Arial", size=10)
BOLD_FONT = Font(name="Arial", size=10, bold=True)
TITLE_FONT = Font(name="Arial", size=14, bold=True)
INPUT_FONT = Font(name="Arial", size=10, color="0000FF")
LINK_FONT = Font(name="Arial", size=10, color="008000")
INPUT_FILL = PatternFill("solid", start_color="FFFF00")
HEADER_FILL = PatternFill("solid", start_color="D9E1F2")


def build_workbook():
    workbook = Workbook()
    write_inputs_sheet(workbook.active)
    write_data_sheet(workbook.create_sheet("Data"))
    write_fields_sheet(workbook.create_sheet("Fields"))
    workbook.save(OUTPUT_PATH)


def write_inputs_sheet(sheet):
    sheet.title = "Inputs"
    write_cell(sheet, "A1", "Bloomberg vol data pull", TITLE_FONT)
    write_cell(sheet, "A3", "Start date", BOLD_FONT)
    write_input(sheet, "B3", START_DATE).number_format = "m/d/yyyy"
    write_cell(sheet, "A4", "End date (blank = today)", BOLD_FONT)
    write_input(sheet, "B4", None).number_format = "m/d/yyyy"
    write_cell(sheet, "A6", "Fields", BOLD_FONT)
    for offset, (field, _) in enumerate(FIELDS):
        write_input(sheet, f"{get_column_letter(2 + offset)}6", field)
    write_legend(sheet)
    write_ticker_table(sheet)
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 30
    sheet.column_dimensions["C"].width = 30
    for column in "DEF":
        sheet.column_dimensions[column].width = 30


def write_legend(sheet):
    write_cell(sheet, "H3", "How to use", BOLD_FONT)
    lines = [
        "1. Type tickers in column A (yellow). Change the yellow key in column B if not a US stock/ETF.",
        "2. Set the start date (B3) and optional end date (B4).",
        "3. Bloomberg tab > Refresh Workbook. Each ticker fills its own block on the Data tab.",
        "4. Blue text on yellow = inputs you edit. Green = links. Fields can be swapped in row 6 (check FLDS <GO>).",
        "Rows 11-13 are example tickers; overwrite or clear them.",
    ]
    for offset, line in enumerate(lines):
        write_cell(sheet, f"H{4 + offset}", line)


def write_ticker_table(sheet):
    headers = ["Ticker", "Yellow key", "Bloomberg security", "Data tab column"]
    for offset, header in enumerate(headers):
        cell = write_cell(sheet, f"{get_column_letter(1 + offset)}10", header, BOLD_FONT)
        cell.fill = HEADER_FILL
    for slot in range(TICKER_SLOTS):
        write_ticker_row(sheet, slot)
    sheet.freeze_panes = "A11"


def write_ticker_row(sheet, slot):
    row = FIRST_TICKER_ROW + slot
    ticker = EXAMPLE_TICKERS[slot] if slot < len(EXAMPLE_TICKERS) else None
    write_input(sheet, f"A{row}", ticker)
    write_input(sheet, f"B{row}", DEFAULT_YELLOW_KEY)
    write_cell(sheet, f"C{row}", f'=IF(A{row}="","",A{row}&" "&B{row})')
    write_cell(sheet, f"D{row}", get_column_letter(block_first_column(slot)))


def write_data_sheet(sheet):
    for slot in range(TICKER_SLOTS):
        write_data_block(sheet, slot)
    sheet.freeze_panes = "A4"


def write_data_block(sheet, slot):
    first_column = block_first_column(slot)
    ticker_row = FIRST_TICKER_ROW + slot
    security_cell = f"Inputs!$C${ticker_row}"
    write_cell(sheet, cell_at(first_column, 1), f"={security_cell}", Font(name="Arial", size=11, bold=True, color="008000"))
    write_block_headers(sheet, first_column)
    write_cell(sheet, cell_at(first_column, 3), bdh_formula(security_cell))
    sheet.column_dimensions[get_column_letter(first_column)].width = 11
    for offset in range(1, len(FIELDS) + 1):
        sheet.column_dimensions[get_column_letter(first_column + offset)].width = 14
    sheet.column_dimensions[get_column_letter(first_column + len(FIELDS) + 1)].width = 3


def write_block_headers(sheet, first_column):
    date_header = write_cell(sheet, cell_at(first_column, 2), "Date", BOLD_FONT)
    date_header.fill = HEADER_FILL
    for offset in range(len(FIELDS)):
        field_cell = f"Inputs!{get_column_letter(2 + offset)}$6"
        header = write_cell(sheet, cell_at(first_column + 1 + offset, 2), f"={field_cell}", LINK_FONT)
        header.fill = HEADER_FILL
        header.alignment = Alignment(wrap_text=True)


def bdh_formula(security_cell):
    last_field_column = get_column_letter(1 + len(FIELDS))
    fields_range = f"Inputs!$B$6:${last_field_column}$6"
    end_date = 'IF(Inputs!$B$4="","",Inputs!$B$4)'
    bdh = f"BDH({security_cell},{fields_range},Inputs!$B$3,{end_date},{BDH_OPTIONS})"
    return f'=IF({security_cell}="","",{bdh})'


def write_fields_sheet(sheet):
    write_cell(sheet, "A1", "Field", BOLD_FONT).fill = HEADER_FILL
    write_cell(sheet, "B1", "What it is", BOLD_FONT).fill = HEADER_FILL
    for offset, (field, meaning) in enumerate(FIELDS):
        write_cell(sheet, f"A{2 + offset}", field)
        write_cell(sheet, f"B{2 + offset}", meaning)
    notes_row = 3 + len(FIELDS)
    write_cell(sheet, f"A{notes_row}", "Notes", BOLD_FONT)
    write_cell(sheet, f"A{notes_row + 1}", "Field names come from Bloomberg user forums; confirm each in FLDS <GO> before backtesting.")
    write_cell(sheet, f"A{notes_row + 2}", "Vol fields are in percentage points (25 = 25%). Divide by 100 before comparing to decimals.")
    write_cell(sheet, f"A{notes_row + 3}", 'BDH options: Dir=V (dates down), Days=T (trading days), Fill=P (carry last value over gaps).')
    write_cell(sheet, f"A{notes_row + 4}", "If a block stays blank after refresh, remove the IF( wrapper so the cell is just =BDH(...).")
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 60


def block_first_column(slot):
    return 1 + slot * BLOCK_WIDTH


def cell_at(column, row):
    return f"{get_column_letter(column)}{row}"


def write_cell(sheet, address, value, font=BASE_FONT):
    cell = sheet[address]
    cell.value = value
    cell.font = font
    return cell


def write_input(sheet, address, value):
    cell = write_cell(sheet, address, value, INPUT_FONT)
    cell.fill = INPUT_FILL
    return cell


if __name__ == "__main__":
    build_workbook()
